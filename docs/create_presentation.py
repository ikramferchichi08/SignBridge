"""Generate the 10-minute SignBridge first-evaluation presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("SignBridge_First_Evaluation.pptx")
NAVY = RGBColor(12, 31, 54)
BLUE = RGBColor(34, 126, 230)
TEAL = RGBColor(30, 184, 166)
ORANGE = RGBColor(244, 157, 59)
LIGHT = RGBColor(241, 247, 252)
MID = RGBColor(84, 105, 126)
WHITE = RGBColor(255, 255, 255)


def textbox(slide, text, x, y, w, h, size=20, color=NAVY, bold=False,
            align=PP_ALIGN.LEFT, font="Aptos", margin=0.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def bullet_box(slide, items, x, y, w, h, size=18, color=NAVY):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.bullet = True
    return shape


def panel(slide, x, y, w, h, fill=WHITE, line=RGBColor(218, 229, 239)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def base_slide(prs, title, number, timing):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = LIGHT
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.18)
    ).fill.solid()
    bar = slide.shapes[-1]
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    textbox(slide, title, 0.55, 0.38, 10.8, 0.55, size=28, bold=True)
    textbox(slide, f"{number:02d}  |  {timing}", 11.25, 0.43, 1.5, 0.35,
            size=11, color=MID, align=PP_ALIGN.RIGHT)
    textbox(slide, "SIGNBRIDGE  •  PROJECT MODULE", 0.6, 7.15, 4.5, 0.22,
            size=9, color=MID, bold=True)
    textbox(slide, str(number), 12.35, 7.12, 0.35, 0.25, size=10,
            color=MID, align=PP_ALIGN.RIGHT)
    return slide


def add_flow_node(slide, label, x, color):
    panel(slide, x, 3.05, 1.65, 1.15, fill=color, line=color)
    textbox(slide, label, x + 0.08, 3.28, 1.49, 0.65, size=16,
            color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(9.6), Inches(-1.1), Inches(5.2), Inches(5.2)
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = TEAL
    slide.shapes[-1].line.fill.background()
    textbox(slide, "SIGNBRIDGE", 0.8, 1.15, 7.5, 0.8, size=44, color=WHITE, bold=True)
    textbox(slide, "Real-time sign language recognition\nfrom video", 0.85, 2.1, 7.6, 1.35,
            size=30, color=RGBColor(220, 239, 250), bold=True)
    textbox(slide, "Project Module  •  First evaluation  •  10-minute presentation",
            0.88, 4.15, 7.4, 0.4, size=17, color=RGBColor(173, 205, 225))
    textbox(slide, "Student A: <name>     |     Student B: <name>",
            0.88, 5.35, 7.4, 0.35, size=16, color=WHITE)
    textbox(slide, "From landmarks to a live transcript", 0.88, 6.65, 6.5, 0.3,
            size=13, color=TEAL, bold=True)

    # 2. Problem and objective
    slide = base_slide(prs, "The problem and our objective", 2, "1:00")
    panel(slide, 0.65, 1.35, 5.75, 4.95)
    textbox(slide, "Communication barriers", 0.95, 1.72, 4.9, 0.45, size=24, color=BLUE, bold=True)
    bullet_box(slide, [
        "Interpreters are rarely available in public services, schools, and workplaces.",
        "Existing solutions can be expensive, difficult to deploy, or not real-time.",
        "Signers need feedback that is understandable and immediate.",
    ], 0.95, 2.35, 4.95, 2.4, size=18)
    textbox(slide, "Goal", 0.95, 5.1, 1.0, 0.35, size=18, color=ORANGE, bold=True)
    textbox(slide, "Recognize isolated signs from a webcam or video and turn them into readable text.", 1.75, 5.0, 4.2, 0.75, size=20, bold=True)
    panel(slide, 6.75, 1.35, 5.9, 4.95, fill=NAVY, line=NAVY)
    textbox(slide, "Project scope", 7.1, 1.72, 4.8, 0.4, size=24, color=TEAL, bold=True)
    bullet_box(slide, [
        "Vocabulary: 50–100 isolated signs",
        "Input: webcam or uploaded video",
        "Output: prediction, confidence, transcript",
        "Target: useful real-time performance on a laptop",
    ], 7.08, 2.35, 4.95, 2.7, size=20, color=WHITE)
    textbox(slide, "Continuous sign-language translation is outside this first scope.", 7.1, 5.45, 4.8, 0.5, size=15, color=RGBColor(203, 224, 238))

    # 3. Functional overview
    slide = base_slide(prs, "How a user interacts with SignBridge", 3, "1:15")
    panel(slide, 0.65, 1.3, 12.0, 4.95)
    steps = [
        ("1", "Open app", "Webcam or video upload"),
        ("2", "See skeleton", "Hands, body, face landmarks"),
        ("3", "Sign", "One isolated sign at a time"),
        ("4", "Get feedback", "Top predictions + confidence"),
        ("5", "Read transcript", "Smoothed running text"),
    ]
    for i, (num, title, desc) in enumerate(steps):
        x = 0.95 + i * 2.3
        panel(slide, x, 2.0, 1.85, 2.85, fill=WHITE)
        slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.55), Inches(2.25), Inches(0.75), Inches(0.75)).fill.solid()
        circle = slide.shapes[-1]
        circle.fill.fore_color.rgb = BLUE if i < 3 else TEAL
        circle.line.fill.background()
        textbox(slide, num, x + 0.55, 2.36, 0.75, 0.35, size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, title, x + 0.12, 3.25, 1.6, 0.35, size=18, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, desc, x + 0.15, 3.82, 1.55, 0.75, size=14, color=MID, align=PP_ALIGN.CENTER)
        if i < 4:
            textbox(slide, "→", x + 1.9, 3.15, 0.35, 0.35, size=25, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "Practice mode: learners can repeat a sign and check whether the system recognizes it correctly.", 1.0, 5.55, 11.0, 0.35, size=17, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    # 4. AI components
    slide = base_slide(prs, "AI components and model comparison", 4, "1:30")
    columns = [
        ("Representation", ["MediaPipe Holistic", "Hands + pose comparison", "Normalize around shoulders", "32–64 frame sequences"], BLUE),
        ("Models", ["Landmarks + GRU", "Landmarks + light Transformer", "CNN + LSTM on frames", "Pretrained video model"], TEAL),
        ("Post-processing", ["Sliding-window buffer", "Majority vote", "Confidence threshold", "Low-confidence warning"], ORANGE),
    ]
    for i, (heading, items, color) in enumerate(columns):
        x = 0.7 + i * 4.2
        panel(slide, x, 1.4, 3.85, 4.8)
        textbox(slide, heading, x + 0.25, 1.72, 3.3, 0.4, size=23, color=color, bold=True)
        bullet_box(slide, items, x + 0.25, 2.35, 3.3, 2.8, size=17)
    textbox(slide, "Why multiple models? We want an honest accuracy–latency trade-off, not only the highest offline score.", 1.0, 6.42, 11.2, 0.35, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 5. Workflow
    slide = base_slide(prs, "Complete workflow: data to deployment", 5, "2:00")
    labels = [("Collect", BLUE), ("Version", TEAL), ("Extract", BLUE), ("Train", ORANGE), ("Compare", TEAL), ("Deploy", BLUE)]
    for i, (label, color) in enumerate(labels):
        add_flow_node(slide, label, 0.65 + i * 2.1, color)
        if i < len(labels) - 1:
            textbox(slide, "→", 2.31 + i * 2.1, 3.38, 0.35, 0.35, size=24, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    panel(slide, 0.85, 4.75, 11.7, 1.1, fill=WHITE)
    textbox(slide, "DVC", 1.15, 5.03, 0.75, 0.45, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "raw videos → landmarks → processed splits → models", 2.05, 5.03, 4.2, 0.45, size=17, bold=True)
    textbox(slide, "MLflow", 6.55, 5.03, 1.0, 0.45, size=20, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "parameters • metrics • confusion matrices • artifacts", 7.7, 5.03, 4.1, 0.45, size=17, bold=True)
    textbox(slide, "Signer-aware split prevents data leakage and gives an honest generalization result.", 1.0, 6.35, 11.2, 0.35, size=17, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    # 6. Evaluation
    slide = base_slide(prs, "Evaluation plan and expected results", 6, "1:30")
    panel(slide, 0.7, 1.35, 5.7, 4.95)
    textbox(slide, "Metrics", 1.0, 1.72, 4.8, 0.4, size=24, color=BLUE, bold=True)
    bullet_box(slide, [
        "Top-1 and top-5 accuracy",
        "Macro-F1 and confusion matrix",
        "FPS, response latency, and model size",
        "Word error rate for short continuous clips",
        "Performance per signer and on custom clips",
    ], 1.0, 2.35, 4.9, 3.15, size=18)
    panel(slide, 6.85, 1.35, 5.75, 4.95, fill=NAVY, line=NAVY)
    textbox(slide, "Targets", 7.18, 1.72, 4.8, 0.4, size=24, color=TEAL, bold=True)
    bullet_box(slide, [
        "60–80% top-1 accuracy is a realistic initial target",
        "Over 20 FPS end-to-end for a lightweight landmark model",
        "Clear accuracy vs speed comparison table",
        "Working live demo plus recorded fallback",
    ], 7.18, 2.35, 4.85, 2.8, size=19, color=WHITE)
    textbox(slide, "We will report limitations, not hide failure cases.", 7.18, 5.5, 4.7, 0.35, size=16, color=RGBColor(203, 224, 238), bold=True)

    # 7. Tools and requirements
    slide = base_slide(prs, "Tools mapped to the professor's requirements", 7, "1:00")
    rows = [
        ("Version control", "GitHub", "Branches, issues, pull requests, visible contributions"),
        ("Data versioning", "DVC", "Reproducible data and model artifacts"),
        ("Experiment tracking", "MLflow", "Parameters, metrics, artifacts, model comparison"),
        ("AI stack", "PyTorch + MediaPipe", "Temporal models and landmark extraction"),
        ("Application", "FastAPI + Streamlit/Gradio", "Free real-time API and interface"),
    ]
    y = 1.45
    for i, (area, tool, purpose) in enumerate(rows):
        panel(slide, 0.8, y, 11.75, 0.82, fill=WHITE if i % 2 == 0 else RGBColor(229, 240, 248))
        textbox(slide, area, 1.05, y + 0.15, 2.3, 0.4, size=16, color=MID, bold=True)
        textbox(slide, tool, 3.45, y + 0.15, 2.4, 0.4, size=18, color=BLUE if i < 3 else TEAL, bold=True)
        textbox(slide, purpose, 5.95, y + 0.15, 6.2, 0.4, size=15, color=NAVY)
        y += 0.9
    textbox(slide, "All selected tools are free/open-source and designed to run on a normal laptop.", 1.0, 6.35, 11.2, 0.35, size=17, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    # 8. Team and progress
    slide = base_slide(prs, "Team roles and continuous progress", 8, "0:45")
    panel(slide, 0.7, 1.4, 5.8, 4.95)
    textbox(slide, "Student A  •  Vision / ML", 1.0, 1.78, 4.9, 0.4, size=23, color=BLUE, bold=True)
    bullet_box(slide, [
        "Dataset preparation and signer-aware splits",
        "Landmark extraction and normalization",
        "GRU, Transformer, CNN+LSTM training",
        "MLflow experiments and comparison",
    ], 1.0, 2.45, 4.95, 2.8, size=18)
    panel(slide, 6.85, 1.4, 5.8, 4.95, fill=WHITE)
    textbox(slide, "Student B  •  Real-time / MLOps / App", 7.15, 1.78, 5.0, 0.4, size=21, color=TEAL, bold=True)
    bullet_box(slide, [
        "Webcam pipeline and temporal smoothing",
        "DVC pipeline and reproducibility",
        "FastAPI + frontend integration",
        "Docker and drift monitoring",
    ], 7.15, 2.45, 4.95, 2.8, size=18)
    textbox(slide, "Both: evaluation, custom clips, report, and presentation", 1.0, 6.45, 11.2, 0.3, size=17, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

    # 9. Closing
    slide = base_slide(prs, "What success looks like", 9, "0:30")
    panel(slide, 1.0, 1.55, 11.3, 3.8, fill=NAVY, line=NAVY)
    textbox(slide, "A useful, measurable, reproducible assistant", 1.45, 2.05, 10.4, 0.55, size=29, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "Recognize isolated signs  •  compare models honestly  •  run in real time", 1.45, 2.95, 10.4, 0.4, size=20, color=RGBColor(206, 232, 244), align=PP_ALIGN.CENTER)
    textbox(slide, "GitHub + DVC + MLflow make the progress visible and reproducible.", 1.45, 3.75, 10.4, 0.4, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "Thank you  |  Questions?", 1.0, 6.15, 11.3, 0.55, size=28, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
